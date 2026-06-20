#!/usr/bin/env python3
"""
Seed daily_drug_demand from an Excel export.

Production demand history should come from ``drug_receipts`` (via upload or
``sync_daily_demand_from_receipts``). This script is for bulk Excel imports only.

Only drug_code, demand_date, and total_quantity are inserted. Extra spreadsheet
columns (drug_id, drug_name, region, facility_name, etc.) are ignored except
when needed to derive missing core fields:
  - drug_code from drug_catalog.id when drug_code is blank
  - total_quantity from daily_demand_units when total_quantity is blank

Rows with the same drug_code + demand_date are aggregated (quantities summed).

Usage (from backend/):
  source bin/activate
  export PYTHONPATH=src
  python scripts/seed_daily_drug_demand.py --xlsx /path/to/daily_drug_demand_100k.xlsx
  python scripts/seed_daily_drug_demand.py --xlsx /path/to/file.xlsx --reset
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from app.core.database import SessionLocal
from app.models.daily_drug_demand import DailyDrugDemand
from app.models.drug_catalog import DrugCatalog

BATCH_SIZE = 2000

CORE_COLUMNS = ("drug_code", "demand_date", "total_quantity")
FALLBACK_COLUMNS = ("drug_id", "daily_demand_units")


def column_index(header: tuple, name: str) -> int:
    try:
        return header.index(name)
    except ValueError as exc:
        raise ValueError(f"Excel missing required column: {name}") from exc


def parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            raise ValueError("empty date")
        try:
            return datetime.fromisoformat(text).date()
        except ValueError:
            return datetime.strptime(text, "%Y-%m-%d").date()
    raise ValueError(f"invalid date: {value!r}")


def parse_quantity(value: object) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError("empty quantity")
    return float(value)


def load_id_to_drug_code(db: Session) -> dict[int, str]:
    return {
        drug_id: drug_code
        for drug_id, drug_code in db.query(DrugCatalog.id, DrugCatalog.drug_code).all()
    }


def load_xlsx_rows(
    xlsx_path: Path,
    *,
    id_to_drug_code: dict[int, str],
) -> tuple[list[dict], int]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook.active

    row_iter = worksheet.iter_rows(values_only=True)
    header = next(row_iter)
    if header is None:
        raise ValueError("Excel file is empty")

    idx_drug_code = column_index(header, "drug_code")
    idx_demand_date = column_index(header, "demand_date")
    idx_total_quantity = column_index(header, "total_quantity")
    idx_drug_id = header.index("drug_id") if "drug_id" in header else None
    idx_daily_demand_units = (
        header.index("daily_demand_units") if "daily_demand_units" in header else None
    )

    aggregated: dict[tuple[str, date], float] = defaultdict(float)
    skipped_rows = 0

    for line_no, row in enumerate(row_iter, start=2):
        try:
            demand_date = parse_date(row[idx_demand_date])

            drug_code = row[idx_drug_code]
            code = str(drug_code).strip() if drug_code else ""
            if not code:
                if idx_drug_id is None or row[idx_drug_id] is None:
                    skipped_rows += 1
                    continue
                code = id_to_drug_code.get(int(row[idx_drug_id]))
                if not code:
                    skipped_rows += 1
                    continue

            quantity_value = row[idx_total_quantity]
            if quantity_value is None and idx_daily_demand_units is not None:
                quantity_value = row[idx_daily_demand_units]
            quantity = parse_quantity(quantity_value)

            aggregated[(code, demand_date)] += quantity
        except (IndexError, TypeError, ValueError):
            skipped_rows += 1
            if line_no <= 5:
                raise ValueError(f"invalid row at line {line_no}") from None

    payloads = [
        {
            "drug_code": drug_code,
            "demand_date": demand_date,
            "total_quantity": round(total_quantity, 4),
        }
        for (drug_code, demand_date), total_quantity in aggregated.items()
    ]
    return payloads, skipped_rows


def reset_demand(db: Session) -> None:
    db.query(DailyDrugDemand).delete(synchronize_session=False)
    db.commit()


def seed_demand_from_xlsx(
    db: Session,
    *,
    xlsx_path: Path,
    reset: bool = False,
) -> tuple[int, int, int]:
    id_to_drug_code = load_id_to_drug_code(db)
    if not id_to_drug_code:
        raise ValueError("drug_catalog is empty; seed drug_catalog before demand data")

    payloads, skipped_rows = load_xlsx_rows(xlsx_path, id_to_drug_code=id_to_drug_code)

    if reset:
        reset_demand(db)
        existing_keys: set[tuple[str, date]] = set()
    else:
        existing_keys = {
            (drug_code, demand_date)
            for drug_code, demand_date in db.query(
                DailyDrugDemand.drug_code,
                DailyDrugDemand.demand_date,
            ).all()
        }

    to_insert = [
        row
        for row in payloads
        if (row["drug_code"], row["demand_date"]) not in existing_keys
    ]
    skipped_existing = len(payloads) - len(to_insert)

    for offset in range(0, len(to_insert), BATCH_SIZE):
        batch = to_insert[offset : offset + BATCH_SIZE]
        db.bulk_insert_mappings(DailyDrugDemand, batch)

    db.commit()
    return len(to_insert), skipped_existing, skipped_rows


def main() -> int:
    parser = argparse.ArgumentParser(description="Seed daily_drug_demand from Excel.")
    parser.add_argument(
        "--xlsx",
        type=Path,
        required=True,
        help="Path to daily_drug_demand Excel file",
    )
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Delete all existing daily_drug_demand rows before seeding",
    )
    args = parser.parse_args()

    if not args.xlsx.is_file():
        print(f"Excel file not found: {args.xlsx}", file=sys.stderr)
        return 1

    db = SessionLocal()
    try:
        inserted, skipped_existing, skipped_rows = seed_demand_from_xlsx(
            db,
            xlsx_path=args.xlsx,
            reset=args.reset,
        )
        total = db.query(DailyDrugDemand).count()
        print(
            f"Inserted {inserted} rows "
            f"(skipped {skipped_existing} existing, "
            f"{skipped_rows} invalid source rows). "
            f"daily_drug_demand now has {total} rows."
        )
        return 0
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
